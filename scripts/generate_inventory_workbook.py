from __future__ import annotations

import ast
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "aervise_source_inventory.xlsx"
AREA_ORDER = [
    "contracts",
    "pipeline",
    "core",
    "interfaces",
    "rendering",
    "observability",
    "config",
    "scripts",
    "tests",
    "other",
]

STAGE_ROWS = [
    {
        "stage": "Stage 0",
        "name": "Entry Point",
        "purpose": "Normalize raw external request into the stable entry envelope.",
        "primary_inputs": "user_id, message, timestamp_utc, location, channel",
        "primary_outputs": "EntryPointEnvelope",
        "contracts": "contracts.intent_schema.UserMessageRequest; EntryPointEnvelope",
        "implementation": "pipeline.intent.parser.InteractionService._normalize_entry_point",
        "api_surface": "POST /interaction/preview; POST /interaction/enrich; POST /interaction/build-payload",
        "notes": "Chat/demo UI currently uses this path as the top-level operator input.",
    },
    {
        "stage": "Stage 1",
        "name": "Intent Recognition",
        "purpose": "Convert raw message into explicit structured intent plus semantic activity profile.",
        "primary_inputs": "EntryPointEnvelope.message",
        "primary_outputs": "IntentRecognitionResult",
        "contracts": "contracts.intent_schema.IntentRecognitionResult; ActivityProfile",
        "implementation": "pipeline.intent.parser.InteractionService._recognize_intent; pipeline.intent.activity_normalizer.ActivityNormalizer",
        "api_surface": "POST /interaction/preview",
        "notes": "Rule-based parser with ontology-backed activity normalization.",
    },
    {
        "stage": "Stage 2",
        "name": "Context Enrichment",
        "purpose": "Fetch and normalize air-quality and weather state.",
        "primary_inputs": "EntryPointEnvelope.location; IntentRecognitionResult",
        "primary_outputs": "EnvironmentState; ContextEnrichmentPreview",
        "contracts": "contracts.environment_schema.EnvironmentState; ContextEnrichmentPreview",
        "implementation": "pipeline.enrichment.service.EnrichmentService; pipeline.enrichment.preview.EnrichmentPreviewService",
        "api_surface": "POST /interaction/enrichment-preview; POST /interaction/enrich",
        "notes": "Uses OpenAQ plus weather providers, with snapshot fallback for resilience.",
    },
    {
        "stage": "Stage 3",
        "name": "Decision Payload Builder",
        "purpose": "Build canonical deterministic-engine input from Stage 0/1/2 outputs.",
        "primary_inputs": "entry_point, intent_recognition, environment_state",
        "primary_outputs": "DecisionInput",
        "contracts": "contracts.decision_input.DecisionInput",
        "implementation": "pipeline.payload_builder.builder.DecisionPayloadBuilder; pipeline.orchestrator.PipelineOrchestrator",
        "api_surface": "POST /interaction/build-payload",
        "notes": "Carries activity_profile forward and derives engine archetypes/defaults.",
    },
    {
        "stage": "Stage 4",
        "name": "Factor Evaluation",
        "purpose": "Compute normalized burdens and penalties from engine input.",
        "primary_inputs": "DecisionInput",
        "primary_outputs": "risk_factors",
        "contracts": "contracts.decision_output.RiskFactorSet",
        "implementation": "core.layers.layer2_factor_eval.DecisionLayer2Evaluator",
        "api_surface": "POST /decision/debug/layers; POST /decision/evaluate; POST /decision/evaluate-rendered",
        "notes": "Deterministic only.",
    },
    {
        "stage": "Stage 5",
        "name": "Policy + Explanation",
        "purpose": "Apply deterministic policy logic and produce structured explanation output.",
        "primary_inputs": "Layer 2 payload",
        "primary_outputs": "DecisionOutput",
        "contracts": "contracts.decision_output.DecisionOutput",
        "implementation": "core.layers.layer3_policy.DecisionLayer3Policy; core.layers.layer4_explanation.DecisionLayer4Explainer",
        "api_surface": "POST /decision/evaluate; POST /decision/debug/layers; POST /decision/evaluate-rendered",
        "notes": "Future-day requests currently degrade conservatively due to forecast limits.",
    },
    {
        "stage": "Stage 6",
        "name": "Rendering",
        "purpose": "Convert structured decision output into stable user-facing messaging.",
        "primary_inputs": "DecisionOutput",
        "primary_outputs": "RenderedDecision; RenderedDecisionEnvelope",
        "contracts": "rendering.formatter.RenderedDecision; RenderedDecisionEnvelope",
        "implementation": "rendering.formatter.DecisionMessageFormatter",
        "api_surface": "POST /decision/render; POST /decision/evaluate-rendered",
        "notes": "Template-based, not LLM-driven.",
    },
    {
        "stage": "Cross-Cut",
        "name": "Observability",
        "purpose": "Persist logs and traces for stage and API execution.",
        "primary_inputs": "Trace id, event payloads, route metadata",
        "primary_outputs": "logs/aervise.log; traces/aervise_trace.jsonl",
        "contracts": "observability.trace_schema; observability.logger",
        "implementation": "observability.runtime; route-level instrumentation",
        "api_surface": "All routed execution paths",
        "notes": "UI propagates a shared trace id across the flow.",
    },
]

API_ROWS = [
    {
        "area": "interaction",
        "method": "POST",
        "path": "/interaction/preview",
        "purpose": "Run Stage 0 and Stage 1 only.",
        "request_contract": "pipeline.intent.schemas.InteractionRequest",
        "response_contract": "pipeline.intent.schemas.InteractionResponse",
        "implementation": "interfaces.api.routes.interaction.preview_interaction",
        "stage_coverage": "0-1",
    },
    {
        "area": "interaction",
        "method": "POST",
        "path": "/interaction/enrichment-preview",
        "purpose": "Preview Stage 2 readiness.",
        "request_contract": "pipeline.intent.schemas.InteractionRequest",
        "response_contract": "contracts.environment_schema.ContextEnrichmentPreview",
        "implementation": "interfaces.api.routes.interaction.preview_enrichment",
        "stage_coverage": "0-2 preview",
    },
    {
        "area": "interaction",
        "method": "POST",
        "path": "/interaction/enrich",
        "purpose": "Run Stage 2 and return normalized environment state.",
        "request_contract": "pipeline.intent.schemas.InteractionRequest",
        "response_contract": "contracts.environment_schema.EnvironmentEnrichmentResponse",
        "implementation": "interfaces.api.routes.interaction.enrich_interaction",
        "stage_coverage": "0-2",
    },
    {
        "area": "interaction",
        "method": "POST",
        "path": "/interaction/build-payload",
        "purpose": "Run Stage 0-3 and return canonical decision input.",
        "request_contract": "pipeline.intent.schemas.InteractionRequest",
        "response_contract": "contracts.decision_input.DecisionInput",
        "implementation": "interfaces.api.routes.interaction.build_decision_payload",
        "stage_coverage": "0-3",
    },
    {
        "area": "decision",
        "method": "POST",
        "path": "/decision/evaluate",
        "purpose": "Run deterministic engine on canonical input.",
        "request_contract": "contracts.decision_input.DecisionInput",
        "response_contract": "contracts.decision_output.DecisionOutput",
        "implementation": "interfaces.api.routes.decision.evaluate_decision",
        "stage_coverage": "4-5",
    },
    {
        "area": "decision",
        "method": "POST",
        "path": "/decision/debug/layers",
        "purpose": "Inspect deterministic layers from canonical input.",
        "request_contract": "contracts.decision_input.DecisionInput",
        "response_contract": "debug payload (dict)",
        "implementation": "interfaces.api.routes.decision.debug_layers",
        "stage_coverage": "3-5 debug",
    },
    {
        "area": "decision",
        "method": "POST",
        "path": "/decision/render",
        "purpose": "Render structured output into stable messaging.",
        "request_contract": "contracts.decision_output.DecisionOutput",
        "response_contract": "rendering.formatter.RenderedDecision",
        "implementation": "interfaces.api.routes.decision.render_decision",
        "stage_coverage": "6",
    },
    {
        "area": "decision",
        "method": "POST",
        "path": "/decision/evaluate-rendered",
        "purpose": "Run engine and rendering in one call.",
        "request_contract": "contracts.decision_input.DecisionInput",
        "response_contract": "rendering.formatter.RenderedDecisionEnvelope",
        "implementation": "interfaces.api.routes.decision.evaluate_and_render",
        "stage_coverage": "4-6",
    },
    {
        "area": "health",
        "method": "GET",
        "path": "/health",
        "purpose": "Basic health/status check.",
        "request_contract": "none",
        "response_contract": "interfaces.api.schemas.HealthResponse",
        "implementation": "interfaces.api.routes.health.health_check",
        "stage_coverage": "n/a",
    },
    {
        "area": "ui",
        "method": "GET",
        "path": "/",
        "purpose": "Operator/demo UI.",
        "request_contract": "none",
        "response_contract": "HTML",
        "implementation": "interfaces.api.routes.ui.index",
        "stage_coverage": "operator surface",
    },
]

DRIFT_ROWS = [
    {
        "category": "Stage 1",
        "item": "Rule-based intent parsing",
        "status": "implemented",
        "source_of_truth": "contracts.intent_schema; pipeline.intent.parser",
        "current_state": "Active production path for Stage 0/1.",
        "gap_or_drift": "",
        "next_action": "Expand ontology over time and keep contracts stable.",
    },
    {
        "category": "Stage 1",
        "item": "Activity semantic normalization",
        "status": "implemented",
        "source_of_truth": "pipeline.intent.activity_normalizer; contracts.intent_schema.ActivityProfile",
        "current_state": "Ontology-backed label + profile output is live.",
        "gap_or_drift": "No embedding fallback yet for unresolved phrases.",
        "next_action": "Add fallback matcher only for unknown activities.",
    },
    {
        "category": "Stage 2",
        "item": "Air + weather enrichment",
        "status": "implemented",
        "source_of_truth": "contracts.environment_schema; pipeline.enrichment.service",
        "current_state": "OpenAQ + weather provider orchestration with snapshot fallback.",
        "gap_or_drift": "Operational quality still depends on provider health.",
        "next_action": "Add stronger provider health and cache semantics.",
    },
    {
        "category": "Stage 2",
        "item": "Future-day weather selection",
        "status": "planned",
        "source_of_truth": "docs/04_interaction_flow_contract.md",
        "current_state": "Future-day requests are parsed, but Stage 2 still returns current/same-day oriented environment state.",
        "gap_or_drift": "Tomorrow/weekend requests are conservatively blocked downstream.",
        "next_action": "Select forecast window by requested horizon/window before Stage 3.",
    },
    {
        "category": "Stage 2",
        "item": "Air-quality forecast",
        "status": "planned",
        "source_of_truth": "docs/01_usecases.md; docs/04_b_environment_schema.md",
        "current_state": "Not implemented.",
        "gap_or_drift": "Best-time and future-day air-aware planning are incomplete.",
        "next_action": "Add AQ forecast source and update forecast capabilities.",
    },
    {
        "category": "Stage 3",
        "item": "Canonical decision input builder",
        "status": "implemented",
        "source_of_truth": "contracts.decision_input; pipeline.payload_builder.builder",
        "current_state": "Strict builder in place with activity profile support.",
        "gap_or_drift": "",
        "next_action": "Keep builder as the only engine entry boundary.",
    },
    {
        "category": "Core",
        "item": "Deterministic factor/policy/explanation engine",
        "status": "implemented",
        "source_of_truth": "core.layers.*; contracts.decision_output",
        "current_state": "Fully deterministic current-state engine path.",
        "gap_or_drift": "Future-day reasoning is intentionally conservative, not fully forecast-driven.",
        "next_action": "Introduce true forecast-aware planning layer only after Stage 2 support exists.",
    },
    {
        "category": "Rendering",
        "item": "Template rendering",
        "status": "implemented",
        "source_of_truth": "rendering.formatter; rendering.templates",
        "current_state": "Stable template-based output layer is active.",
        "gap_or_drift": "No LLM rendering path.",
        "next_action": "Keep templates primary; consider optional LLM only later.",
    },
    {
        "category": "Interfaces",
        "item": "Operator/demo UI",
        "status": "implemented",
        "source_of_truth": "interfaces.api.templates.index.html",
        "current_state": "Chat-style top view plus stage inspector.",
        "gap_or_drift": "",
        "next_action": "Keep frontend minimal and debugging-oriented.",
    },
    {
        "category": "Observability",
        "item": "Logs and traces",
        "status": "implemented",
        "source_of_truth": "observability.runtime; route instrumentation",
        "current_state": "Trace-id-based logging and trace persistence are active.",
        "gap_or_drift": "No dedicated trace exploration UI beyond the copyable trace id.",
        "next_action": "Add trace lookup or filtered log inspection if needed.",
    },
    {
        "category": "Legacy",
        "item": "Old direct raw-request engine path",
        "status": "deprecated",
        "source_of_truth": "README.md; docs/04_*",
        "current_state": "Removed from active API flow.",
        "gap_or_drift": "Legacy docs still exist in parts of docs/02 and docs/03.",
        "next_action": "Continue pruning or marking old docs explicitly.",
    },
]


@dataclass
class SymbolRow:
    area: str
    module: str
    file_path: str
    line: int
    symbol_type: str
    class_name: str
    symbol_name: str
    signature_or_members: str
    decorators_or_bases: str
    notes: str = ""


@dataclass
class ClassContext:
    name: str
    lineno: int
    kind: str
    bases: list[str] = field(default_factory=list)


class InventoryVisitor(ast.NodeVisitor):
    def __init__(self, module: str, file_path: str, area: str) -> None:
        self.module = module
        self.file_path = file_path
        self.area = area
        self.rows: list[SymbolRow] = []
        self.class_stack: list[ClassContext] = []

    def visit_ClassDef(self, node: ast.ClassDef) -> None:
        bases = [self._expr_to_text(base) for base in node.bases]
        decorators = [self._expr_to_text(item) for item in node.decorator_list]
        enum_members = self._enum_members(node)
        class_kind = self._class_kind(node=node, bases=bases, decorators=decorators)
        signature = ", ".join(enum_members) if enum_members else ""
        self.rows.append(
            SymbolRow(
                area=self.area,
                module=self.module,
                file_path=self.file_path,
                line=node.lineno,
                symbol_type=class_kind,
                class_name="",
                symbol_name=node.name,
                signature_or_members=signature,
                decorators_or_bases=", ".join(decorators or bases),
            )
        )
        if enum_members:
            for member in enum_members:
                self.rows.append(
                    SymbolRow(
                        area=self.area,
                        module=self.module,
                        file_path=self.file_path,
                        line=node.lineno,
                        symbol_type="enum_member",
                        class_name=node.name,
                        symbol_name=member,
                        signature_or_members="",
                        decorators_or_bases="",
                    )
                )

        self.class_stack.append(ClassContext(name=node.name, lineno=node.lineno, kind=class_kind, bases=bases))
        self.generic_visit(node)
        self.class_stack.pop()

    def visit_FunctionDef(self, node: ast.FunctionDef) -> None:
        self._record_function(node)

    def visit_AsyncFunctionDef(self, node: ast.AsyncFunctionDef) -> None:
        self._record_function(node, is_async=True)

    def _record_function(self, node: ast.FunctionDef | ast.AsyncFunctionDef, is_async: bool = False) -> None:
        decorators = [self._expr_to_text(item) for item in node.decorator_list]
        args = [arg.arg for arg in node.args.args]
        if node.args.vararg:
            args.append(f"*{node.args.vararg.arg}")
        args.extend(arg.arg for arg in node.args.kwonlyargs)
        if node.args.kwarg:
            args.append(f"**{node.args.kwarg.arg}")
        signature = f"{'async ' if is_async else ''}{node.name}({', '.join(args)})"
        if self.class_stack:
            owner = self.class_stack[-1]
            symbol_type = "method"
            if "staticmethod" in decorators:
                symbol_type = "staticmethod"
            elif "classmethod" in decorators:
                symbol_type = "classmethod"
            self.rows.append(
                SymbolRow(
                    area=self.area,
                    module=self.module,
                    file_path=self.file_path,
                    line=node.lineno,
                    symbol_type=symbol_type,
                    class_name=owner.name,
                    symbol_name=node.name,
                    signature_or_members=signature,
                    decorators_or_bases=", ".join(decorators),
                )
            )
        else:
            self.rows.append(
                SymbolRow(
                    area=self.area,
                    module=self.module,
                    file_path=self.file_path,
                    line=node.lineno,
                    symbol_type="function",
                    class_name="",
                    symbol_name=node.name,
                    signature_or_members=signature,
                    decorators_or_bases=", ".join(decorators),
                )
            )
        self.generic_visit(node)

    def _class_kind(self, *, node: ast.ClassDef, bases: list[str], decorators: list[str]) -> str:
        if any(base.endswith("Enum") for base in bases):
            return "enum"
        if any(base.endswith("BaseModel") for base in bases):
            return "schema_model"
        if any(base.endswith("APIRouter") for base in bases):
            return "router_class"
        if "dataclass" in decorators:
            return "dataclass"
        return "class"

    @staticmethod
    def _enum_members(node: ast.ClassDef) -> list[str]:
        members: list[str] = []
        for item in node.body:
            if isinstance(item, ast.Assign):
                for target in item.targets:
                    if isinstance(target, ast.Name):
                        members.append(target.id)
            elif isinstance(item, ast.AnnAssign) and isinstance(item.target, ast.Name):
                members.append(item.target.id)
        return members

    @staticmethod
    def _expr_to_text(node: ast.AST) -> str:
        try:
            return ast.unparse(node)
        except Exception:
            return type(node).__name__


def iter_python_files(root: Path) -> Iterable[Path]:
    for path in sorted(root.rglob("*.py")):
        if any(part.startswith(".") for part in path.parts):
            continue
        yield path


def area_for(path: Path) -> str:
    top = path.parts[0]
    return top if top in AREA_ORDER else "other"


def module_name(path: Path) -> str:
    relative = path.with_suffix("")
    return ".".join(relative.parts)


def collect_inventory() -> list[SymbolRow]:
    rows: list[SymbolRow] = []
    for path in iter_python_files(ROOT):
        relative = path.relative_to(ROOT)
        area = area_for(relative)
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(relative))
        visitor = InventoryVisitor(module=module_name(relative), file_path=str(relative), area=area)
        visitor.visit(tree)
        rows.extend(visitor.rows)
    return rows


def write_sheet(workbook: Workbook, title: str, rows: list[SymbolRow]) -> None:
    ws = workbook.create_sheet(title=title)
    headers = [
        "Area",
        "Module",
        "File",
        "Line",
        "Type",
        "Class",
        "Name",
        "Signature / Members",
        "Decorators / Bases",
        "Notes",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append(
            [
                row.area,
                row.module,
                row.file_path,
                row.line,
                row.symbol_type,
                row.class_name,
                row.symbol_name,
                row.signature_or_members,
                row.decorators_or_bases,
                row.notes,
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    width_targets = {
        1: 16,
        2: 42,
        3: 42,
        4: 8,
        5: 18,
        6: 24,
        7: 30,
        8: 50,
        9: 36,
        10: 30,
    }
    for index, width in width_targets.items():
        ws.column_dimensions[get_column_letter(index)].width = width


def write_dict_sheet(workbook: Workbook, title: str, rows: list[dict[str, str]], headers: list[str]) -> None:
    ws = workbook.create_sheet(title=title)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row in rows:
        ws.append([row.get(header.lower().replace(" ", "_"), "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 28 if index < 5 else 40


def build_workbook(rows: list[SymbolRow]) -> Workbook:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"

    overview["A1"] = "Aervise Source Inventory"
    overview["A1"].font = Font(size=14, bold=True)
    overview["A3"] = "Generated File"
    overview["B3"] = str(OUTPUT_PATH.relative_to(ROOT))
    overview["A4"] = "Tracked Symbols"
    overview["B4"] = len(rows)
    overview["A5"] = "Tracked Areas"
    overview["B5"] = len({row.area for row in rows})
    overview["A7"] = "Area"
    overview["B7"] = "Symbol Count"
    overview["C7"] = "Files"
    for cell in overview[7]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    area_rows = []
    for area in AREA_ORDER:
        scoped = [row for row in rows if row.area == area]
        if not scoped:
            continue
        area_rows.append((area, len(scoped), len({row.file_path for row in scoped})))
    for index, (area, symbol_count, file_count) in enumerate(area_rows, start=8):
        overview[f"A{index}"] = area
        overview[f"B{index}"] = symbol_count
        overview[f"C{index}"] = file_count

    overview["E3"] = "Tracking Notes"
    overview["E3"].font = Font(bold=True)
    overview["E4"] = "Use the per-area sheets as the generated inventory."
    overview["E5"] = "Use the Notes column on those sheets for manual ownership/status comments."
    overview["E6"] = "Regenerate with: python scripts/generate_inventory_workbook.py"
    overview["E7"] = "Use the Stage Map and API Contracts sheets as the generated architecture tracker."
    overview.column_dimensions["A"].width = 20
    overview.column_dimensions["B"].width = 16
    overview.column_dimensions["C"].width = 12
    overview.column_dimensions["E"].width = 70

    for area in AREA_ORDER:
        scoped = [row for row in rows if row.area == area]
        if scoped:
            write_sheet(workbook, area.title(), scoped)

    write_dict_sheet(
        workbook,
        "Stage Map",
        STAGE_ROWS,
        ["Stage", "Name", "Purpose", "Primary Inputs", "Primary Outputs", "Contracts", "Implementation", "API Surface", "Notes"],
    )
    write_dict_sheet(
        workbook,
        "API Contracts",
        API_ROWS,
        ["Area", "Method", "Path", "Purpose", "Request Contract", "Response Contract", "Implementation", "Stage Coverage"],
    )
    write_dict_sheet(
        workbook,
        "Drift TODO",
        DRIFT_ROWS,
        ["Category", "Item", "Status", "Source Of Truth", "Current State", "Gap Or Drift", "Next Action"],
    )
    write_sheet(workbook, "All Symbols", rows)
    return workbook


def main() -> None:
    rows = collect_inventory()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(rows)
    workbook.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
